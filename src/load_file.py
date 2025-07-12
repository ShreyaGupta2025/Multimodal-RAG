import os
import json
import requests
import streamlit as st
import pandas as pd
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from bs4 import BeautifulSoup
from typing import List

def load_file(file_path: str = None, url: str = None):
    if url:
        documents = load_from_url(url)
    elif file_path:
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            documents = load_pdf(file_path)
        elif file_extension == '.txt':
            documents = load_txt(file_path)
        elif file_extension == '.json':
            documents = load_json(file_path)
        elif file_extension == '.html':
            documents = load_html(file_path)
        elif file_extension == '.docx':
            documents = load_docx(file_path)
        elif file_extension in ['.xls', '.xlsx']:
            documents = load_excel(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")
    else:
        raise ValueError("Either file_path or url must be provided.")
    return documents
    
def load_pdf(file_path: str) -> List[str]:
    with open(file_path, "rb") as f:
        reader = PdfReader(f)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
    return text.splitlines()

def load_txt(file_path: str) -> List[str]:
    with open(file_path, "r") as f:
        return f.readlines()

def load_json(file_path: str) -> List[str]:
    with open(file_path, "r") as f:
        data = json.load(f)
    return [json.dumps(data)]

def load_html(file_path: str) -> List[str]:
    with open(file_path, "r") as f:
        soup = BeautifulSoup(f, "html.parser")
        return [soup.get_text()]

def load_docx(file_path: str) -> List[str]:
    doc = Document(file_path)
    return [para.text for para in doc.paragraphs if para.text]

def load_excel(file_path: str) -> List[str]:
    df = pd.read_excel(file_path)
    return df.astype(str).apply(lambda x: ' '.join(x), axis=1).tolist()

def load_from_url(url: str) -> List[str]:
    response = requests.get(url)
    if response.status_code == 200:
        soup = BeautifulSoup(response.content, "html.parser")
        return [soup.get_text()]
    else:
        raise ValueError(f"Failed to fetch URL: {response.status_code}")

def read_file(file_path):
    # Read content from a file.
    try:
        with open(file_path, 'r') as file:
            return file.read()
    except FileNotFoundError:
        print(f"The file '{file_path}' does not exist.")
        return ""
    except IOError:
        print("An error occurred while reading the file.")
        return ""

def write_to_file(file_path, content):
    # Write content to a file.
    try:
        file_name = file_path + '_sample.txt'
        with open(file_name, 'w') as file:
            file.write(content)
        st.write(f"Content successfully written to {file_name}")
    except IOError:
        st.write("An error occurred while writing to the file.") 



